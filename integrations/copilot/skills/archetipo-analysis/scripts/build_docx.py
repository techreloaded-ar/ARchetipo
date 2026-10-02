#!/usr/bin/env python3
"""
build_docx.py — turns the Markdown produced from the analysis templates into a Word document.

Usage:
    python build_docx.py analisi.md --images screens --out Analisi-Funzionale.docx

Supported Markdown (the subset used by template-documento.md and template-sezione.md):
    # Title                       -> Title (first level-1 heading) ; later level-1 headings -> Heading 1
    ## / ### / #### / #####       -> Heading 1 / 2 / 3 / 4
    paragraph line                -> Normal (every line is its own paragraph; no soft wrapping)
    - item / * item               -> List Bullet (indent of 2+ spaces -> List Bullet 2)
    1. item                       -> List Number
    | a | b |  (+ |---|---| row)   -> Word table, header row bold and repeated on every page
    **bold** *italic* _italic_ `code`   inside any paragraph, list item or table cell
    ![caption](images/file.png)   -> inline picture 16 cm wide + "Figura N — caption" (style Caption)
    > quote                       -> Normal, italic
    ---                           -> ignored (template separator)
    ``` fenced block ```          -> monospace paragraphs

The script never invents content. A missing image leaves a visible placeholder in the document
and is listed in the report. The report is printed on stdout; the exit code is 2 when the document
has residual Markdown markers or missing images, 0 otherwise. The file is written in both cases.

Requires python-docx only.
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from dataclasses import dataclass, field

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

IMAGE_WIDTH_CM = 16.0
TABLE_FONT_PT = 9
BODY_FONT_PT = 10.5

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*"  # bold
    r"|`[^`\n]+`"  # code
    r"|(?<![\w*])\*(?!\s)[^*\n]+?(?<!\s)\*(?![\w*])"  # *italic*
    r"|(?<![\w_])_(?!\s)[^_\n]+?(?<!\s)_(?![\w_]))"  # _italic_
)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBER_RE = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
HR_RE = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")


@dataclass
class Report:
    headings: dict = field(default_factory=dict)
    paragraphs: int = 0
    list_items: int = 0
    tables: int = 0
    table_rows: int = 0
    images_ok: list = field(default_factory=list)
    images_missing: list = field(default_factory=list)
    bold_runs: int = 0
    italic_runs: int = 0
    code_runs: int = 0
    residual: list = field(default_factory=list)


# --------------------------------------------------------------------------- inline Markdown


def add_inline(paragraph, text: str, report: Report, base_bold=False, base_italic=False, size_pt=None):
    """Append runs to `paragraph`, interpreting **bold**, *italic*, _italic_ and `code`."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos : m.start()], base_bold, base_italic, False, size_pt)
        token = m.group(0)
        if token.startswith("**"):
            _run(paragraph, token[2:-2], True, base_italic, False, size_pt)
            report.bold_runs += 1
        elif token.startswith("`"):
            _run(paragraph, token[1:-1], base_bold, base_italic, True, size_pt)
            report.code_runs += 1
        else:
            _run(paragraph, token[1:-1], base_bold, True, False, size_pt)
            report.italic_runs += 1
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], base_bold, base_italic, False, size_pt)


def _run(paragraph, text, bold, italic, code, size_pt):
    if text == "":
        return
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if code:
        run.font.name = "Consolas"
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), "Consolas")
        rfonts.set(qn("w:hAnsi"), "Consolas")
    if size_pt:
        run.font.size = Pt(size_pt)


# --------------------------------------------------------------------------- block helpers


def split_cells(line: str) -> list[str]:
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|") and not body.endswith("\\|"):
        body = body[:-1]
    cells, cur, i = [], [], 0
    while i < len(body):
        ch = body[i]
        if ch == "\\" and i + 1 < len(body) and body[i + 1] == "|":
            cur.append("|")
            i += 2
            continue
        if ch == "|":
            cells.append("".join(cur).strip())
            cur = []
        else:
            cur.append(ch)
        i += 1
    cells.append("".join(cur).strip())
    return cells


def set_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def shade_cell(cell, hex_fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def add_table(doc, header: list[str], rows: list[list[str]], report: Report):
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, text in enumerate(header):
        cell = hdr.cells[i]
        shade_cell(cell, "E5F3F4")
        p = cell.paragraphs[0]
        add_inline(p, text, report, base_bold=True, size_pt=TABLE_FONT_PT)
    for r in rows:
        r = (r + [""] * ncols)[:ncols]
        cells = table.add_row().cells
        for i, text in enumerate(r):
            add_inline(cells[i].paragraphs[0], text, report, size_pt=TABLE_FONT_PT)
    report.tables += 1
    report.table_rows += len(rows)
    doc.add_paragraph()  # breathing space after a table


def add_image(doc, caption: str, path: str, images_dir: str, figure_no: int, report: Report):
    candidates = [path, os.path.join(images_dir, path), os.path.join(images_dir, os.path.basename(path))]
    found = next((c for c in candidates if os.path.isfile(c)), None)
    if found:
        doc.add_picture(found, width=Cm(IMAGE_WIDTH_CM))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        report.images_ok.append(found)
    else:
        p = doc.add_paragraph()
        run = p.add_run(f"[immagine non prodotta: {path}]")
        run.italic = True
        run.font.color.rgb = RGBColor(0x97, 0x24, 0x38)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        report.images_missing.append(path)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(cap, f"Figura {figure_no} — {caption}", report)


# --------------------------------------------------------------------------- main conversion


def convert(md_text: str, images_dir: str, out_path: str) -> Report:
    report = Report()
    doc = Document()
    _page_setup(doc)

    lines = md_text.splitlines()
    i, n = 0, len(lines)
    title_done = False
    figure_no = 0
    in_fence = False

    while i < n:
        line = lines[i]

        if line.strip().startswith("```"):
            in_fence = not in_fence
            i += 1
            continue
        if in_fence:
            p = doc.add_paragraph()
            _run(p, line, False, False, True, 9)
            i += 1
            continue

        if not line.strip() or HR_RE.match(line):
            i += 1
            continue

        m = HEADING_RE.match(line)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1 and not title_done:
                p = doc.add_paragraph(style="Title")
                title_done = True
                key = "Title"
            else:
                lvl = max(1, min(level - 1 if title_done else level, 9))
                p = doc.add_paragraph(style=f"Heading {lvl}")
                key = f"Heading {lvl}"
            add_inline(p, text, report)
            report.headings[key] = report.headings.get(key, 0) + 1
            i += 1
            continue

        m = IMAGE_RE.match(line.strip())
        if m:
            figure_no += 1
            add_image(doc, m.group(1).strip(), m.group(2).strip(), images_dir, figure_no, report)
            i += 1
            continue

        if line.lstrip().startswith("|") and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]):
            header = split_cells(line)
            rows = []
            i += 2
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(split_cells(lines[i]))
                i += 1
            add_table(doc, header, rows, report)
            continue

        m = BULLET_RE.match(line)
        if m:
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_inline(p, m.group(2), report)
            report.list_items += 1
            i += 1
            continue

        m = NUMBER_RE.match(line)
        if m:
            style = "List Number 2" if len(m.group(1)) >= 2 else "List Number"
            p = doc.add_paragraph(style=style)
            add_inline(p, m.group(2), report)
            report.list_items += 1
            i += 1
            continue

        if line.lstrip().startswith(">"):
            p = doc.add_paragraph()
            add_inline(p, line.lstrip()[1:].strip(), report, base_italic=True)
            report.paragraphs += 1
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, line.strip(), report)
        report.paragraphs += 1
        i += 1

    _scan_residual(doc, report)
    doc.save(out_path)
    return report


def _page_setup(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(BODY_FONT_PT)
    normal.paragraph_format.space_after = Pt(4)


RESIDUAL_RE = re.compile(r"(\*\*|(?<!\w)`|^#{1,6}\s|!\[.*?\]\(|\]\()", re.M)


def _scan_residual(doc, report: Report):
    def check(text, where):
        if RESIDUAL_RE.search(text):
            report.residual.append(f"{where}: {text[:80]}")

    for p in doc.paragraphs:
        if p.style.name.startswith("Caption"):
            continue
        check(p.text, "paragraph")
    for t_idx, t in enumerate(doc.tables):
        for r_idx, r in enumerate(t.rows):
            for c in r.cells:
                check(c.text, f"table {t_idx + 1} row {r_idx + 1}")


def print_report(report: Report, out_path: str):
    size = os.path.getsize(out_path) if os.path.exists(out_path) else 0
    print("== build_docx report ==")
    print(f"output: {out_path} ({size / 1024:.0f} KB)")
    print("headings: " + ", ".join(f"{k}={v}" for k, v in sorted(report.headings.items())))
    print(f"paragraphs: {report.paragraphs} · list items: {report.list_items}")
    print(f"tables: {report.tables} · table rows: {report.table_rows}")
    print(f"inline runs: bold={report.bold_runs} italic={report.italic_runs} code={report.code_runs}")
    print(f"images inserted: {len(report.images_ok)}")
    print(f"images missing: {len(report.images_missing)}" + ("" if not report.images_missing else " -> " + ", ".join(report.images_missing)))
    print(f"residual markdown markers: {len(report.residual)}")
    for r in report.residual[:10]:
        print("   " + r)
    ok = not report.images_missing and not report.residual
    print("RESULT: " + ("OK" if ok else "CHECK FAILED (see lines above)"))
    return ok


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass
    ap = argparse.ArgumentParser(description="Markdown (analysis templates) -> Word")
    ap.add_argument("markdown")
    ap.add_argument("--images", default="screens", help="folder with the PNG files referenced by the Markdown")
    ap.add_argument("--out", default="Analisi-Funzionale.docx")
    args = ap.parse_args(argv)
    with open(args.markdown, encoding="utf-8") as fh:
        md_text = fh.read()
    report = convert(md_text, args.images, args.out)
    ok = print_report(report, args.out)
    return 0 if ok else 2


if __name__ == "__main__":
    sys.exit(main())
